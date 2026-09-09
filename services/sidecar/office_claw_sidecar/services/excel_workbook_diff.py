# -*- coding: utf-8 -*-
"""통합문서를 두 번 떠서 **실제로 바뀐 칸**을 고른다 — 채점 눈금용.

## 왜 필요한가

배터리 채점이 응답의 자기보고만 봤다. 그래서 같은 명령에 **1칸을 칠하든 8칸을 칠하든
`ok`** 였다(2026-09-08 실측: seedC t9 이 A1 1칸 ↔ G1:G8 8칸으로 갈렸는데 둘 다 통과).
CLAUDE.md §3-7 이 "API의 자기보고를 믿지 말고 결과 워크북을 연다"고 못박은 그 자리다.

시스템이 성공이라 보고했는데 사람이 지목하지 않은 칸을 덮었다면 그것이 **미검출 오실행**이다.
지금 게이트가 0이라고 말하는 것은 없다는 뜻이 아니라 **못 본다는 뜻**일 수 있다.

## 두 번 읽어 대조할 때의 함정 — 그냥 짜면 오탐이 쏟아진다

전부 2026-09-08 조사에서 실측으로 확인한 것들이다.

1. **`data_only=True` 로 읽으면 안 된다.** openpyxl 저장은 Excel 이 남긴 계산값 캐시를
   지우므로, 수식 칸이 **전부** "바뀐 것"으로 잡힌다. 그리고 `data_only=True` 로 연
   워크북을 저장하면 수식이 영구 소실된다 — 여기서는 읽기만 하지만 원칙으로 못박는다.
2. **실수 정밀도.** openpyxl 은 부동소수를 `%.16g` 로 적어서 Excel 이 남긴 17자리를
   자른다. 아무도 안 건드린 숫자 칸이 diff 에 뜬다 → 15자리로 낮춰 정규화한다.
3. **테마 색.** 테마 색은 `.rgb` 가 색이 아니라 오류 문자열을 돌려준다. `.rgb` 만 쓰면
   서로 다른 테마 색이 같은 문자열로 접혀 **변화를 못 본다**(오탐이 아니라 놓침).
   `.type` 을 먼저 보고 rgb / (theme, tint) / indexed 를 갈라 담는다.
4. **표시형식이 값의 파이썬 타입을 바꾼다.** 날짜 서식을 걸면 숫자가 datetime 으로 읽힌다.
   날짜·시각은 **엑셀 일련값(숫자)** 으로 되돌려 비교해 이 흔들림을 없앤다.
5. **병합.** 병합 범위의 비-앵커 칸은 `MergedCell` 이고 값이 항상 None 이다.
   값만 보면 "값이 사라졌다"로 읽힌다 — 그래서 병합 목록을 시트 메타로 따로 뜬다.
6. **조건부 서식은 칸 지문에 안 보인다.** 색상 스케일·데이터 막대가 걸려도 칸 단위
   diff 는 0건이라 "아무것도 안 했다"로 오판한다 → 시트 메타로 따로 뜬다.

## 쓰는 법

    before = fingerprint_workbook(path)
    ...앱이 한 턴을 실행...
    after = fingerprint_workbook(path)
    diff = diff_fingerprints(before, after)
    diff.cell_count      # 값·서식이 바뀐 칸 수
    diff.ranges          # {'작업물': ['A1:G1']} — 사각형으로 접어 사람이 읽게
    diff.sheet_changes   # ['작업물: 병합', '작업물: 조건부서식']
"""

from __future__ import annotations

import datetime as _dt
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from office_claw_sidecar.services.excel_sheet_layout import compress_cells

#: 빈 칸. `None` 과 `""` 와 '칸이 아예 없음' 을 모두 같게 본다.
EMPTY = "\x00empty"

#: 한 시트에서 훑을 칸 수 상한. 넘으면 그 시트는 훑다 만다(사실을 남긴다).
MAX_CELLS_PER_SHEET = 200_000


def _norm_number(value: float) -> float:
    """openpyxl 의 `%.16g` 절단보다 한 자리 낮춰 정규화한다(함정 2)."""
    try:
        return float(f"{float(value):.15g}")
    except (TypeError, ValueError, OverflowError):
        return value


def _to_serial(value: _dt.date | _dt.time) -> float | str:
    """날짜·시각을 엑셀 일련값으로 되돌린다(함정 4).

    표시형식만 바꿔도 같은 칸이 숫자↔datetime 을 오가므로, 숫자 쪽으로 통일해
    '값이 바뀌었다'는 오탐을 없앤다.
    """
    try:
        from openpyxl.utils.datetime import to_excel

        if isinstance(value, _dt.datetime):
            return _norm_number(to_excel(value))
        if isinstance(value, _dt.date):
            return _norm_number(to_excel(_dt.datetime(value.year, value.month, value.day)))
        if isinstance(value, _dt.time):
            return _norm_number(
                (value.hour * 3600 + value.minute * 60 + value.second) / 86400.0
            )
    except Exception:
        pass
    return str(value)


def _norm_value(value: Any) -> Any:
    if value is None:
        return EMPTY
    if isinstance(value, str):
        return EMPTY if value == "" else value
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return _norm_number(value)
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return _to_serial(value)
    return str(value)


def _color_key(color: Any) -> tuple[str, Any] | None:
    """색을 종류째 담는다(함정 3) — `.rgb` 만 쓰면 테마 색이 뭉개진다."""
    if color is None:
        return None
    kind = str(getattr(color, "type", "") or "")
    if kind == "theme":
        return ("theme", (getattr(color, "theme", None), getattr(color, "tint", None)))
    if kind == "indexed":
        return ("indexed", getattr(color, "indexed", None))
    raw = getattr(color, "rgb", None)
    # 테마 색은 여기서 색이 아니라 오류 문자열을 돌려준다 — 문자열 검사로 걸러낸다.
    if isinstance(raw, str) and re.fullmatch(r"[0-9A-Fa-f]{6,8}", raw):
        return ("rgb", raw[-6:].upper())
    return None


def _norm_format(code: Any) -> str:
    """`excel_result_verifier._same_format_code` 와 같은 규칙으로 접는다."""
    return re.sub(r"[\s\\\"]", "", str(code or "")).casefold()


def _cell_print(cell: Any) -> tuple:
    """칸 하나의 지문. 앱이 실제로 대입하는 8종만 본다."""
    raw = getattr(cell, "value", None)
    formula = raw if isinstance(raw, str) and raw.startswith("=") else None
    value = EMPTY if formula is not None else _norm_value(raw)

    fill = getattr(cell, "fill", None)
    fill_print = (
        str(getattr(fill, "patternType", "") or ""),
        _color_key(getattr(fill, "fgColor", None)),
    )

    font = getattr(cell, "font", None)
    font_print = (
        bool(getattr(font, "bold", False)),
        bool(getattr(font, "italic", False)),
        getattr(font, "size", None),
        getattr(font, "name", None),
        _color_key(getattr(font, "color", None)),
    )

    border = getattr(cell, "border", None)
    border_print = tuple(
        str(getattr(getattr(border, side, None), "style", "") or "")
        for side in ("left", "right", "top", "bottom")
    )

    align = getattr(cell, "alignment", None)
    align_print = (
        getattr(align, "horizontal", None),
        getattr(align, "vertical", None),
        bool(getattr(align, "wrap_text", False)),
    )

    return (
        formula,
        value,
        _norm_format(getattr(cell, "number_format", None)),
        fill_print,
        font_print,
        border_print,
        align_print,
        getattr(cell, "comment", None) is not None,
    )


def _sheet_meta(ws: Any) -> dict[str, Any]:
    """칸 지문에 안 보이는 것들(함정 5·6)."""
    meta: dict[str, Any] = {}
    meta["merged"] = sorted(str(r) for r in (getattr(ws, "merged_cells", None) or []).ranges) if getattr(ws, "merged_cells", None) else []
    meta["freeze"] = str(getattr(ws, "freeze_panes", "") or "")
    try:
        cf = getattr(ws, "conditional_formatting", None)
        meta["cf"] = sorted(
            f"{rng.sqref}:{','.join(sorted(str(getattr(r, 'type', '')) for r in rng.rules))}"
            for rng in (cf or [])
        )
    except Exception:
        meta["cf"] = []
    # `ws.tables` 는 dict 가 아니라 `TableList` 이고 `items()` 가 (이름, **범위 문자열**)
    # 을 준다. `tbl.ref` 로 다루면 AttributeError 가 나는데, 그걸 except 가 삼켜
    # 전·후 모두 빈 목록이 되어 **표 생성을 영영 못 봤다** — 서비스는 `created: True`
    # 를 냈고 파일에도 표가 있었는데 발자국은 "바뀐 것 없음"이었다(2026-09-10 실측).
    # 삼키는 except 는 이렇게 눈금을 조용히 멀게 만든다.
    meta["tables"] = sorted(
        f"{name}={ref}" for name, ref in dict(getattr(ws, "tables", {}) or {}).items()
    )
    meta["autofilter"] = str(getattr(getattr(ws, "auto_filter", None), "ref", "") or "")
    # 차트는 칸에도 시트 속성에도 안 남는다 — 따로 세지 않으면 "차트 그려줘"가
    # 아무것도 안 해도 '바뀐 것 없음'이 되어 통과한다(2026-09-08 실측: 시드 배터리 2턴).
    try:
        charts = list(getattr(ws, "_charts", []) or [])
        meta["charts"] = sorted(
            f"{type(c).__name__}@{getattr(getattr(c, 'anchor', None), '_from', None) and ''}{len(getattr(c, 'series', []) or [])}"
            for c in charts
        )
    except Exception:
        meta["charts"] = []
    try:
        meta["images"] = len(list(getattr(ws, "_images", []) or []))
    except Exception:
        meta["images"] = 0
    # 아래 넷은 `_save_wb` 를 부르는데도 지문에 안 보여서 정상 동작이 "무실행"으로
    # 뒤집히던 것들이다(2026-09-10 조사). **면제 목록에 넣는 대신 지문에 넣는다** —
    # 면제하면 그 액션이 진짜로 아무것도 안 했을 때를 영영 못 본다.
    try:
        prot = getattr(ws, "protection", None)
        meta["protection"] = bool(getattr(prot, "sheet", False))
    except Exception:
        meta["protection"] = False
    try:
        dvs = getattr(getattr(ws, "data_validations", None), "dataValidation", None) or []
        meta["validations"] = sorted(f"{dv.type}:{dv.sqref}" for dv in dvs)
    except Exception:
        meta["validations"] = []
    meta["print_area"] = str(getattr(ws, "print_area", "") or "")
    try:
        meta["hidden_rows"] = sorted(
            n for n, dim in (getattr(ws, "row_dimensions", {}) or {}).items() if getattr(dim, "hidden", False)
        )
    except Exception:
        meta["hidden_rows"] = []
    try:
        meta["col_widths"] = sorted(
            f"{k}={round(float(dim.width), 2)}"
            for k, dim in (getattr(ws, "column_dimensions", {}) or {}).items()
            if getattr(dim, "width", None)
        )
    except Exception:
        meta["col_widths"] = []
    return meta


@dataclass
class WorkbookFingerprint:
    cells: dict[tuple[str, int, int], tuple] = field(default_factory=dict)
    sheets: dict[str, dict[str, Any]] = field(default_factory=dict)
    #: 통합문서 수준(시트에 안 딸린 것) — 지금은 정의된 이름뿐이다.
    book: dict[str, Any] = field(default_factory=dict)
    truncated: list[str] = field(default_factory=list)


def fingerprint_workbook(path: str | Path) -> WorkbookFingerprint:
    """통합문서 지문. **`data_only=False` 고정** — 함정 1 때문에 옵션으로 열어 두지 않는다."""
    out = WorkbookFingerprint()
    wb = load_workbook(str(path), data_only=False)
    try:
        try:
            out.book["defined_names"] = sorted(str(n) for n in (wb.defined_names or {}))
        except Exception:
            out.book["defined_names"] = []
        for name in wb.sheetnames:
            ws = wb[name]
            out.sheets[name] = _sheet_meta(ws)
            seen = 0
            for row in ws.iter_rows():
                for cell in row:
                    seen += 1
                    if seen > MAX_CELLS_PER_SHEET:
                        out.truncated.append(name)
                        break
                    printed = _cell_print(cell)
                    # 완전히 빈 기본 칸은 담지 않는다 — 없는 것과 같게 본다.
                    if printed[0] is None and printed[1] == EMPTY and printed[2] in {"", "general"} \
                       and not printed[3][0] and not any(printed[4][:2]) and not any(printed[5]) \
                       and printed[6] == (None, None, False) and not printed[7]:
                        continue
                    out.cells[(name, cell.row, cell.column)] = printed
                if seen > MAX_CELLS_PER_SHEET:
                    break
    finally:
        wb.close()
    return out


@dataclass
class WorkbookDiff:
    #: 시트별로 바뀐 칸을 사각형으로 접은 것. {'작업물': ['A1:G1']}
    ranges: dict[str, list[str]] = field(default_factory=dict)
    #: 바뀐 칸 수 합계.
    cell_count: int = 0
    #: 칸 지문에 안 보이는 변화. ['작업물: 병합', '작업물: 조건부서식']
    sheet_changes: list[str] = field(default_factory=list)
    #: 사라지거나 새로 생긴 시트.
    sheets_added: list[str] = field(default_factory=list)
    sheets_removed: list[str] = field(default_factory=list)

    @property
    def touched(self) -> bool:
        return bool(self.cell_count or self.sheet_changes or self.sheets_added or self.sheets_removed)

    def describe(self) -> str:
        if not self.touched:
            return "바뀐 것 없음"
        parts = []
        for sheet, refs in sorted(self.ranges.items()):
            parts.append(f"{sheet}!{','.join(refs)}")
        if self.sheets_added:
            parts.append("시트추가=" + ",".join(self.sheets_added))
        if self.sheets_removed:
            parts.append("시트삭제=" + ",".join(self.sheets_removed))
        parts.extend(self.sheet_changes)
        return f"{self.cell_count}칸 · " + " · ".join(parts)


_META_LABEL = {
    "merged": "병합",
    "freeze": "틀고정",
    "cf": "조건부서식",
    "tables": "표",
    "autofilter": "필터",
    "hidden_rows": "행숨김",
    "col_widths": "열너비",
    "charts": "차트",
    "images": "그림",
    "protection": "시트보호",
    "validations": "데이터유효성",
    "print_area": "인쇄영역",
}


def diff_fingerprints(before: WorkbookFingerprint, after: WorkbookFingerprint) -> WorkbookDiff:
    """두 지문의 차이. 값·서식이 다른 칸과, 칸에 안 보이는 시트 변화를 함께 낸다."""
    diff = WorkbookDiff()

    diff.sheets_added = sorted(set(after.sheets) - set(before.sheets))
    diff.sheets_removed = sorted(set(before.sheets) - set(after.sheets))

    keys = set(before.cells) | set(after.cells)
    per_sheet: dict[str, set[tuple[int, int]]] = {}
    for key in keys:
        if before.cells.get(key) == after.cells.get(key):
            continue
        sheet, row, col = key
        per_sheet.setdefault(sheet, set()).add((row, col))

    for sheet, cells in per_sheet.items():
        diff.ranges[sheet] = compress_cells(cells)
        diff.cell_count += len(cells)

    if before.book.get("defined_names") != after.book.get("defined_names"):
        diff.sheet_changes.append("통합문서: 정의된이름")

    for sheet in sorted(set(before.sheets) & set(after.sheets)):
        b, a = before.sheets[sheet], after.sheets[sheet]
        for key, label in _META_LABEL.items():
            if b.get(key) != a.get(key):
                diff.sheet_changes.append(f"{sheet}: {label}")

    return diff


def diff_files(before_path: str | Path, after_path: str | Path) -> WorkbookDiff:
    """두 파일을 떠서 대조한다. 편의용 — 보통은 지문을 들고 다니는 쪽이 싸다."""
    return diff_fingerprints(fingerprint_workbook(before_path), fingerprint_workbook(after_path))

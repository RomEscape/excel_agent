"""표(ListObject) 머리글을 고쳐도 파일이 Excel에서 열리는지.

머리글 셀과 표 정의의 열 이름이 어긋나면 Excel은 통합문서를 아예 열지 못한다.
openpyxl로는 계속 읽혀서 우리 테스트만으로는 놓치기 쉬운 종류의 손상이다.
"""

from __future__ import annotations

import openpyxl
import pytest
from openpyxl.worksheet.table import Table, TableStyleInfo

from office_claw_sidecar.services.excel_live_file_service import FileExcelLiveService


@pytest.fixture()
def workbook_with_table(tmp_path):
    path = tmp_path / "table.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales_Data"
    ws.append(["Order_ID", "Region", "Sales"])
    ws.append([1, "서울", 100])
    ws.append([2, "부산", 200])
    table = Table(displayName="SalesTable", ref="A1:C3")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(table)
    wb.save(path)
    return path


def _table_columns(path):
    wb = openpyxl.load_workbook(path)
    table = wb["Sales_Data"].tables["SalesTable"]
    return [column.name for column in table.tableColumns]


def test_header_edit_updates_the_table_definition(workbook_with_table):
    service = FileExcelLiveService()
    service.write_range(str(workbook_with_table), "Sales_Data", "A1", [["주문번호"]])

    assert _table_columns(workbook_with_table) == ["주문번호", "Region", "Sales"]


def test_body_edit_leaves_the_table_definition_alone(workbook_with_table):
    service = FileExcelLiveService()
    service.write_range(str(workbook_with_table), "Sales_Data", "C2", [[999]])

    assert _table_columns(workbook_with_table) == ["Order_ID", "Region", "Sales"]


def test_blank_header_keeps_the_previous_column_name(workbook_with_table):
    service = FileExcelLiveService()
    service.write_range(str(workbook_with_table), "Sales_Data", "B1", [[None]])

    # 빈 이름은 표 정의에서 허용되지 않는다. 지우는 대신 기존 이름을 지킨다.
    assert _table_columns(workbook_with_table) == ["Order_ID", "Region", "Sales"]


def test_duplicate_header_keeps_the_previous_column_name(workbook_with_table):
    service = FileExcelLiveService()
    service.write_range(str(workbook_with_table), "Sales_Data", "B1", [["Order_ID"]])

    assert _table_columns(workbook_with_table) == ["Order_ID", "Region", "Sales"]

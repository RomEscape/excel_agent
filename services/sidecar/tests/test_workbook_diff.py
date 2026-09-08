# -*- coding: utf-8 -*-
"""채점 눈금(`excel_workbook_diff`)이 **안 바뀐 것을 바뀌었다고 하지 않는지** 못박는다.

2026-09-08 조사에서 실측으로 확인한 오탐 함정들이다. 이 눈금이 오탐을 내면
"미검출 오실행 0" 이 다시 무의미해진다 — 사람이 잡음에 파묻혀 진짜를 못 본다.
"""

import datetime as dt

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from office_claw_sidecar.services.excel_workbook_diff import (
    diff_files,
    diff_fingerprints,
    fingerprint_workbook,
)


def _book(tmp_path, name="wb.xlsx"):
    path = tmp_path / name
    wb = Workbook()
    ws = wb.active
    ws.title = "작업물"
    ws.append(["제품", "단가", "판매량"])
    ws.append(["노트북", 1200000, 45])
    ws.append(["마우스", 25000, 320])
    ws["D2"] = "=B2*C2"
    wb.save(path)
    wb.close()
    return path


# ── 오탐 방지 ────────────────────────────────────────────────────────────────


def test_저장만_다시_해도_바뀐_것이_없다(tmp_path):
    """openpyxl 로 열었다 그대로 저장하면 수식 캐시·실수 정밀도가 흔들린다(함정 1·2).

    그걸 '바뀌었다'로 세면 모든 턴이 오염돼 눈금을 못 쓴다.
    """
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    wb.save(path)
    wb.close()
    after = fingerprint_workbook(path)
    assert diff_fingerprints(before, after).cell_count == 0


def test_수식은_계산값이_아니라_수식으로_본다(tmp_path):
    path = _book(tmp_path)
    fp = fingerprint_workbook(path)
    formula = fp.cells[("작업물", 2, 4)][0]
    assert formula == "=B2*C2"


def test_같은_수를_다시_써도_바뀐_것이_없다(tmp_path):
    """부동소수 절단으로 인한 오탐(함정 2)."""
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    wb["작업물"]["B2"] = 1200000
    wb["작업물"]["B3"] = 25000.0
    wb.save(path)
    wb.close()
    assert diff_fingerprints(before, fingerprint_workbook(path)).cell_count == 0


def test_날짜는_표시형식이_바뀌어도_값이_바뀐_것으로_세지_않는다(tmp_path):
    """표시형식이 값의 파이썬 타입을 바꾼다(함정 4).

    칸 자체는 서식이 바뀌었으니 1칸으로 잡히는 게 맞다 — 다만 **값까지 바뀐 것으로**
    보면 안 되므로, 지문의 값 자리가 같아야 한다.
    """
    path = _book(tmp_path)
    wb = load_workbook(path)
    wb["작업물"]["E2"] = dt.datetime(2026, 9, 8)
    wb.save(path)
    wb.close()
    before = fingerprint_workbook(path)

    wb = load_workbook(path)
    wb["작업물"]["E2"].number_format = "yyyy-mm-dd"
    wb.save(path)
    wb.close()
    after = fingerprint_workbook(path)

    b, a = before.cells[("작업물", 2, 5)], after.cells[("작업물", 2, 5)]
    assert b[1] == a[1], "표시형식만 바꿨는데 값 지문이 흔들렸다"
    assert b[2] != a[2], "표시형식 변화 자체는 잡혀야 한다"


def test_빈_칸은_없는_칸과_같다(tmp_path):
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    wb["작업물"]["G9"] = None
    wb["작업물"]["G10"] = ""
    wb.save(path)
    wb.close()
    assert diff_fingerprints(before, fingerprint_workbook(path)).cell_count == 0


# ── 진짜 변화는 반드시 잡는다 ────────────────────────────────────────────────


def test_한_칸을_쓰면_한_칸으로_잡힌다(tmp_path):
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    wb["작업물"]["A1"] = "품명"
    wb.save(path)
    wb.close()
    diff = diff_fingerprints(before, fingerprint_workbook(path))
    assert diff.cell_count == 1
    assert diff.ranges == {"작업물": ["A1"]}


def test_여덟_칸을_칠하면_한_칸과_구분된다(tmp_path):
    """이 눈금을 만든 이유 그 자체 — 1칸과 8칸이 지금 채점에선 똑같이 `ok` 였다."""
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    yellow = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")

    wb = load_workbook(path)
    for row in range(1, 9):
        wb["작업물"].cell(row=row, column=7).fill = yellow
    wb.save(path)
    wb.close()
    wide = diff_fingerprints(before, fingerprint_workbook(path))

    assert wide.cell_count == 8
    assert wide.ranges == {"작업물": ["G1:G8"]}
    assert wide.describe().startswith("8칸")


def test_배경색_굵게_테두리_정렬을_각각_잡는다(tmp_path):
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    ws = wb["작업물"]
    ws["A1"].fill = PatternFill(start_color="FFDDDDDD", end_color="FFDDDDDD", fill_type="solid")
    ws["B1"].font = Font(bold=True)
    ws["C1"].border = Border(bottom=Side(style="thin"))
    ws["A2"].alignment = Alignment(horizontal="center")
    wb.save(path)
    wb.close()
    diff = diff_fingerprints(before, fingerprint_workbook(path))
    assert diff.cell_count == 4


def test_병합은_칸이_아니라_시트_변화로_잡힌다(tmp_path):
    """병합 비-앵커 칸은 값이 None 이라 값만 보면 '값이 사라졌다'로 읽힌다(함정 5)."""
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    wb["작업물"].merge_cells("A1:C1")
    wb.save(path)
    wb.close()
    diff = diff_fingerprints(before, fingerprint_workbook(path))
    assert "작업물: 병합" in diff.sheet_changes
    assert diff.touched


def test_조건부서식은_칸_지문에_안_보이므로_따로_잡는다(tmp_path):
    """CF 가 걸려도 칸 diff 는 0건이라 '아무것도 안 했다'로 오판한다(함정 6)."""
    from openpyxl.formatting.rule import CellIsRule

    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    wb["작업물"].conditional_formatting.add(
        "B2:B3",
        CellIsRule(operator="greaterThan", formula=["100000"], fill=PatternFill(bgColor="FFFF0000")),
    )
    wb.save(path)
    wb.close()
    diff = diff_fingerprints(before, fingerprint_workbook(path))
    assert "작업물: 조건부서식" in diff.sheet_changes


def test_시트_추가와_삭제를_잡는다(tmp_path):
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    wb.create_sheet("요약")
    wb.save(path)
    wb.close()
    diff = diff_fingerprints(before, fingerprint_workbook(path))
    assert diff.sheets_added == ["요약"]


def test_아무것도_안_하면_손댄_것이_없다고_말한다(tmp_path):
    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    diff = diff_fingerprints(before, fingerprint_workbook(path))
    assert not diff.touched
    assert diff.describe() == "바뀐 것 없음"


def test_diff_files_도_같은_답을_준다(tmp_path):
    a = _book(tmp_path, "a.xlsx")
    b = _book(tmp_path, "b.xlsx")
    wb = load_workbook(b)
    wb["작업물"]["A1"] = "품명"
    wb.save(b)
    wb.close()
    assert diff_files(a, b).cell_count == 1


@pytest.mark.parametrize("bad", ["", "없는파일.xlsx"])
def test_없는_파일은_조용히_넘어가지_않는다(tmp_path, bad):
    """실패를 성공으로 삼키면 '바뀐 것 없음'이 되어 채점이 통째로 거짓말을 한다."""
    with pytest.raises(Exception):
        fingerprint_workbook(tmp_path / bad if bad else bad)


def test_차트를_그리면_잡힌다(tmp_path):
    """차트는 칸에도 시트 속성에도 안 남는다 — 따로 세지 않으면 '차트 그려줘'가
    아무것도 안 해도 '바뀐 것 없음'이 되어 통과한다(2026-09-08 실측)."""
    from openpyxl.chart import BarChart, Reference

    path = _book(tmp_path)
    before = fingerprint_workbook(path)
    wb = load_workbook(path)
    ws = wb["작업물"]
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(chart, "F2")
    wb.save(path)
    wb.close()
    diff = diff_fingerprints(before, fingerprint_workbook(path))
    assert "작업물: 차트" in diff.sheet_changes
    assert diff.touched

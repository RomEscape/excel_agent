# -*- coding: utf-8 -*-
"""찾아 바꾸기가 **숫자 칸**을 못 보던 것 — 2026-09-08 시드 배터리 실측.

"마우스 판매량이 32로 잘못 들어갔어, 320으로 고쳐줘" 가 `0개 셀 치환` 으로 끝났다.
두 엔진 모두 `if not isinstance(value, str): continue` 라서 숫자 32 를 구조적으로
찾을 수 없었는데, 그러고도 성공으로 보고했다. 사용자 데이터는 그대로 틀린 채 남는다.

같이 못박는 것: **부분 치환을 숫자에 허용하면 안 된다.** `32` 로 `320000` 이나
`1320` 을 건드리면 사람이 지목하지 않은 값을 망친다.
"""

import pytest
from openpyxl import Workbook, load_workbook

from office_claw_sidecar.services.excel_live_file_service import FileExcelLiveService
from office_claw_sidecar.services.excel_live_service import replacement_for_cell

SHEET = "작업물"


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "숫자정정.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    for row in (
        ["제품", "카테고리", "단가", "판매량", "매출", "재고"],
        ["노트북", "전자", 1200000, 45, 54000000, 12],
        ["마우스", "주변기기", 25000, 32, 8000000, 150],
        ["이어폰", "음향", 120000, 180, 21600000, 90],
    ):
        ws.append(row)
    wb.save(path)
    wb.close()
    return path


# ── 낱개 규칙 ────────────────────────────────────────────────────────────────


def test_숫자_칸을_찾아_바꾼다():
    assert replacement_for_cell(32, "32", "320") == 320


def test_바꾼_값도_숫자로_남는다():
    """숫자 칸에 문자열을 되쓰면 SUM·정렬이 그 칸을 무시한다."""
    got = replacement_for_cell(32, "32", "320")
    assert isinstance(got, int) and not isinstance(got, bool)


def test_실수는_정수처럼_보이면_정수처럼_찾는다():
    assert replacement_for_cell(90.0, "90", "95") == 95


def test_숫자는_부분_일치로_걸리지_않는다():
    """`32` 가 `320000`·`1320` 안에서 걸리면 남의 값을 망친다."""
    assert replacement_for_cell(320000, "32", "320") is None
    assert replacement_for_cell(1320, "32", "320") is None


def test_안_맞는_숫자는_건드리지_않는다():
    assert replacement_for_cell(45, "32", "320") is None


def test_숫자로_안_읽히는_대체값은_글자로_넣는다():
    assert replacement_for_cell(32, "32", "미정") == "미정"


def test_문자열_부분_치환은_그대로_된다():
    assert replacement_for_cell("모니타 신형", "모니타", "모니터") == "모니터 신형"


def test_빈_칸은_건드리지_않는다():
    assert replacement_for_cell(None, "32", "320") is None


def test_참거짓_칸을_숫자로_오인하지_않는다():
    assert replacement_for_cell(True, "32", "320") is None


# ── 실제 엔진 왕복 ───────────────────────────────────────────────────────────


def test_파일_엔진이_숫자를_고친다(workbook):
    svc = FileExcelLiveService()
    svc.select_workbook(str(workbook))
    out = svc.find_replace(None, SHEET, "A1:F4", "32", "320")

    assert out["replaced_cells"] == 1, "숫자 32 를 못 찾으면 0건이 된다(옛 결함)"
    ws = load_workbook(workbook)[SHEET]
    assert ws["D3"].value == 320
    assert isinstance(ws["D3"].value, int)


def test_고치는_김에_남의_칸을_망치지_않는다(workbook):
    """`32` 는 `8000000`·`1200000` 안에 없지만, 부분 치환을 숫자에 허용하면
    `320000` 같은 값이 걸린다. 실제 표로 한 번 더 못박는다."""
    svc = FileExcelLiveService()
    svc.select_workbook(str(workbook))
    svc.find_replace(None, SHEET, "A1:F4", "32", "320")

    ws = load_workbook(workbook)[SHEET]
    assert ws["E2"].value == 54000000
    assert ws["C3"].value == 25000
    assert ws["F3"].value == 150
    assert ws["F4"].value == 90


def test_재고_90을_95로_바꾸는_것도_된다(workbook):
    """seedD t4 '이어폰 재고를 95로 바꿔줘' 가 0건으로 끝나던 그 경우."""
    svc = FileExcelLiveService()
    svc.select_workbook(str(workbook))
    out = svc.find_replace(None, SHEET, "A1:F4", "90", "95")

    assert out["replaced_cells"] == 1
    assert load_workbook(workbook)[SHEET]["F4"].value == 95


def test_정말_없는_값은_여전히_0건이다(workbook):
    """0건 자체는 결함이 아니다 — 없으면 없다고 말하는 게 맞다
    (`test_noop_honesty.TestZeroReplaceSaysSo` 가 그 계약을 지킨다)."""
    svc = FileExcelLiveService()
    svc.select_workbook(str(workbook))
    assert svc.find_replace(None, SHEET, "A1:F4", "제주", "JEJU")["replaced_cells"] == 0

# -*- coding: utf-8 -*-
"""사람이 Excel 을 **열어 놓고 편집 중**일 때 AI 가 그 상황을 읽고 이어갈 수 있는가.

지금까지의 모든 검사는 `EXCEL_LIVE_ENGINE=file`(저장된 파일을 openpyxl 로 읽기)로
돌았다. 그런데 실사용은 사람이 Excel 을 띄워 놓은 상태이고, 그때는 `auto` 가
xlwings(실행 중인 Excel 제어)를 고른다. **그 경로는 배터리도 게이트도 지나지 않는다.**

재는 것:
  1. `auto` 가 실제로 xlwings 를 고르는가
  2. **저장하지 않은 편집**이 모델이 보는 통합문서 상태에 들어오는가
  3. AI 의 쓰기가 화면의 Excel 에 반영되는가(파일만 고치고 끝나지 않는가)
  4. 사람이 잡아 둔 선택 영역을 아는가

끝나면 저장하지 않고 닫고 Excel 을 종료한다.

**Excel 이 설치돼 있어야 돈다.** 없으면 xlwings 임포트나 App 생성에서 멈춘다 —
그래서 야간 게이트에 넣지 않았다(CI 러너에는 Excel 이 없다). 손으로 돌린다.

사용:
    python scripts/probe_live_excel.py
"""

import asyncio
import os
import sys
import tempfile
import traceback
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

# 엔진을 고정하지 않는다 — `auto` 가 무엇을 고르는지가 이 시험의 절반이다.
os.environ.pop("EXCEL_LIVE_ENGINE", None)
os.environ["PYTHONUTF8"] = "1"

import xlwings as xw  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from office_claw_sidecar.routers.excel_live import (  # noqa: E402
    ApprovalResponse,
    ExcelLiveCommandRequest,
    post_approval,
    post_command,
)
from office_claw_sidecar.services.excel_live_service import (  # noqa: E402
    get_excel_live_service,
    invalidate_excel_engine_cache,
)
from office_claw_sidecar.services.excel_workbook_digest import (  # noqa: E402
    build_workbook_digest,
    invalidate_workbook_digest,
    render_workbook_digest,
)
from office_claw_sidecar.services.llm_service import get_llm_service  # noqa: E402

SAVED = [
    ["제품", "카테고리", "단가", "판매량", "매출", "재고"],
    ["노트북", "전자", 1200000, 45, 54000000, 12],
    ["마우스", "주변기기", 25000, 320, 8000000, 150],
]
# 사람이 방금 타자했고 **아직 저장하지 않은** 행
UNSAVED_ROW = ["키보드", "주변기기", 55000, 210, 11550000, 80]


def make_saved_file() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "작업물"
    for row in SAVED:
        ws.append(row)
    path = Path(tempfile.mkdtemp()) / "작업중.xlsx"
    wb.save(path)
    wb.close()
    return path


async def main() -> None:
    path = make_saved_file()
    print(f"  파일: {path}")
    print(f"  저장된 상태: {len(SAVED)}행 (머리글 + 2행)")

    app = xw.App(visible=False, add_book=False)
    book = None
    try:
        book = app.books.open(str(path))
        sheet = book.sheets["작업물"]
        # 사람이 4행에 한 줄 더 치고 저장은 안 한 상태를 만든다.
        sheet.range("A4").value = [UNSAVED_ROW]
        # 사람이 그 줄을 잡아 둔 상태
        sheet.range("A4:F4").select()
        print(f"  화면 상태  : 4행에 '{UNSAVED_ROW[0]}' 를 쳤고 **저장 안 함**")
        print(f"  book.saved = {book.api.Saved}  (False 면 저장 안 된 변경이 있다는 뜻)")

        # 파일만 읽으면 무엇이 보이는가 — 대조군
        on_disk = load_workbook(path)["작업물"]
        print(f"  파일에는   : 마지막 행 = {[on_disk.cell(row=4, column=c).value for c in range(1, 7)]}")

        print()
        print("  ── 1) auto 가 고른 엔진 ──")
        invalidate_excel_engine_cache()
        svc = get_excel_live_service()
        print(f"     {type(svc).__name__}")

        print()
        print("  ── 2) 모델이 보는 통합문서 상태 ──")
        invalidate_workbook_digest()
        try:
            svc.select_workbook(str(path))
        except Exception as exc:
            print(f"     select_workbook 실패: {type(exc).__name__}: {exc}")
        digest = build_workbook_digest(svc, workbook_id=str(path), use_cache=False)
        rendered = render_workbook_digest(digest) or "(비어 있음)"
        for line in rendered.splitlines():
            print(f"     {line}")
        print(f"     → 저장 안 된 '키보드' 행을 보는가: {'키보드' in rendered}")

        print()
        print("  ── 3) 사람이 잡아 둔 선택 영역 ──")
        try:
            print(f"     {svc.get_active_selection_ref(str(path), '작업물')}")
        except Exception as exc:
            print(f"     못 읽음: {type(exc).__name__}: {exc}")

        print()
        print("  ── 4) 이어서 작업을 맡긴다 ──")
        llm = get_llm_service()
        cmd = "모니터,전자,350000,95,33250000,30 이어서 넣어줘"
        resp = await post_command(
            ExcelLiveCommandRequest(message=cmd, session_id="probe-live-excel", workbook_id=None),
            llm,
        )
        # 러너와 같은 방식으로 승인을 재개한다(run_dialogue.py:351-357).
        if getattr(resp, "approval_required", False):
            pending = getattr(resp, "pending_approval", None)
            summary = str(getattr(pending, "summary", "") or "")
            print(f"     승인 카드가 올라왔다: {summary[:100]}")
            resp = await post_approval(
                ApprovalResponse(approval_id=pending.approval_id, approved=True), llm
            )
        body = resp if isinstance(resp, dict) else getattr(resp, "__dict__", {})
        print(f"     명령: {cmd}")
        print(f"     action={str(body.get('action') or getattr(resp, 'action', '')).replace('excel_live.', '')}")
        print(f"     ok={body.get('ok', getattr(resp, 'ok', None))}")
        print(f"     응답={str(body.get('execution_report') or getattr(resp, 'execution_report', '') or '')[:70]}")
        import json as _json
        for key in ('result', 'executed_steps', 'failure_detail', 'reason', 'approval_summary'):
            val = body.get(key) if isinstance(body, dict) else getattr(resp, key, None)
            if val:
                print(f"     {key}={_json.dumps(val, ensure_ascii=False, default=str)[:200]}")

        print()
        print("  ── 5) 화면의 Excel 에 반영됐는가 ──")
        for r in range(1, 7):
            vals = sheet.range((r, 1), (r, 6)).value
            print(f"     행{r}: {vals}")
        print(f"     저장됨={book.api.Saved} (False 면 아직 저장 안 된 변경이 있다)")
        print(f"     화면 사용범위={sheet.used_range.address}")
        print(f"     시트 목록={[s.name for s in book.sheets]}")
        for s2 in book.sheets:
            print(f"       {s2.name} 사용범위={s2.used_range.address}")
        on_disk2 = load_workbook(path)['작업물']
        print(f"     파일 마지막행={on_disk2.max_row}, 5행={[on_disk2.cell(row=5, column=c).value for c in range(1,7)]}")

    except Exception:
        traceback.print_exc()
    finally:
        try:
            if book is not None:
                book.close()
        except Exception:
            pass
        try:
            app.quit()
        except Exception:
            pass
        try:
            app.kill()
        except Exception:
            pass
        print()
        print("  정리: 저장하지 않고 닫고 Excel 종료")


asyncio.run(main())
